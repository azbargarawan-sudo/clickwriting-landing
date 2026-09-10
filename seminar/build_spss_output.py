# -*- coding: utf-8 -*-
"""SPSS-style output document computed from the SIMULATED dataset (N=65). Also writes .sav and .sps files."""
import numpy as np, pandas as pd, datetime
from scipy import stats
import statsmodels.api as sm
from statsmodels.stats.diagnostic import lilliefors
import pyreadstat
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

df = pd.read_csv('simulated_data_N65_for_JASP.csv')
N = len(df)
STAMP = datetime.datetime.now().strftime('%d-%b-%Y %H:%M:%S').upper()

VLAB = {
 'Group_LD': 'קבוצה (לקות למידה)', 'Gender': 'מין', 'Age': 'גיל', 'Grade': 'כיתה', 'Sector': 'מגזר',
 'ADHD': 'אבחנת הפרעת קשב', 'Accommodations': 'התאמות בלמידה/בחינות', 'Medication': 'טיפול תרופתי לקשב',
 'Screen_Leisure_Avg': 'זמן מסך פנאי יומי ממוצע (שעות)', 'Screen_Leisure_WD': 'זמן מסך פנאי ביום חול (שעות)',
 'Screen_Leisure_WE': 'זמן מסך פנאי בסוף שבוע (שעות)', 'Screen_Bedroom': 'מסך בחדר השינה', 'Screen_BeforeBed': 'מסך לפני השינה',
 'Bedtime_WD': 'שעת שינה ביום חול', 'Sleep_Dur_WD': 'משך שינה ביום חול (שעות)', 'Sleep_Dur_WE': 'משך שינה בסוף שבוע (שעות)',
 'Weekend_Shift': 'איחור שעת השינה בסוף השבוע (שעות)', 'Sleep_Onset': 'זמן הירדמות', 'Sleepiness_Total': 'ישנוניות יומית (סכום 8 פריטים, 0-32)',
 'PA_Days60': 'ימים בשבוע עם 60 דקות פעילות גופנית', 'PA_Hours_Vig': 'שעות פעילות נמרצת בשבוע', 'PA_Club': 'חוג/קבוצת ספורט',
 'Grade_Avg': 'ממוצע ציונים', 'Grade_Math': 'ציון מתמטיקה', 'Grade_English': 'ציון אנגלית', 'Func_Total': 'תפקוד יומי בלמידה (ממוצע 5 פריטים, 1-5)',
 'Meet_Screen': 'עמידה בהמלצת מסכים', 'Meet_Sleep': 'עמידה בהמלצת שינה', 'Meet_PA': 'עמידה בהמלצת פעילות גופנית', 'Guidelines_Count': 'מספר המלצות שמתקיימות',
 'Guidelines_Grp': 'מספר המלצות (מקובץ)',
 'S1': 'נרדם/ישנוני בשיעור', 'S2': 'נרדם/ישנוני בשיעורי בית', 'S3': 'ערני רוב היום (הפוך)', 'S3R': 'ערני רוב היום (מהופך)', 'S4': 'עייף ועצבני במהלך היום',
 'S5': 'קשה לקום בבוקר', 'S6': 'חוזר לישון אחרי שהתעורר', 'S7': 'צריך שיעירו אותו', 'S8': 'מרגיש שצריך יותר שינה',
 'F1': 'מצליח להתרכז בשיעורים', 'F2': 'מכין שיעורי בית בזמן', 'F3': 'מגיע בזמן לשיעור הראשון', 'F4': 'מרגיש עייף בשיעורים (הפוך)', 'F4R': 'מרגיש עייף בשיעורים (מהופך)', 'F5': 'מאמין ביכולת להצליח',
}
VALLAB = {
 'Group_LD': {0: 'ללא לקות למידה', 1: 'עם לקות למידה'}, 'Gender': {1: 'בן', 2: 'בת'},
 'Grade': {7: "ז'", 8: "ח'", 9: "ט'", 10: "י'", 11: 'י"א', 12: 'י"ב'}, 'Sector': {1: 'יהודי', 2: 'ערבי', 3: 'דרוזי'},
 'ADHD': {0: 'לא', 1: 'כן'}, 'Accommodations': {0: 'לא', 1: 'כן'}, 'Medication': {0: 'לא', 1: 'כן'},
 'Meet_Screen': {0: 'לא עומד', 1: 'עומד'}, 'Meet_Sleep': {0: 'לא עומד', 1: 'עומד'}, 'Meet_PA': {0: 'לא עומד', 1: 'עומד'},
 'Guidelines_Count': {0: '0', 1: '1', 2: '2', 3: '3'}, 'Guidelines_Grp': {0: 'אף המלצה', 1: 'המלצה אחת', 2: 'שתיים או שלוש'},
 'Screen_Bedroom': {0: 'לא', 1: 'כן'}, 'PA_Club': {0: 'לא', 1: 'לא תחרותי', 2: 'תחרותי'},
}
df['S3R'] = 4 - df['S3']; df['F4R'] = 6 - df['F4']
df['Guidelines_Grp'] = df['Guidelines_Count'].clip(upper=2)

# ---------- number formatting like SPSS ----------
def f(x, d=3):
    if x is None or (isinstance(x, float) and np.isnan(x)): return ''
    s = f"{x:.{d}f}"
    if abs(x) < 1 and d > 0: s = s.replace('0.', '.', 1) if s.startswith('0.') else s.replace('-0.', '-.', 1)
    return s
def fp(p):
    return '<.001' if p < .001 else f(p, 3)
def stars(p): return '**' if p < .01 else ('*' if p < .05 else '')

# ---------- document helpers ----------
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); st.element.rPr.rFonts.set(qn('w:cs'), 'Arial')
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)

def para(text, bold=False, size=10, color=None, align=None, italic=False, space_after=4):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.name = 'Arial'
    if color: r.font.color.rgb = RGBColor(*color)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p
def proc_title(t): para(t, bold=True, size=14, space_after=6)
def tbl_title(t): para(t, bold=True, size=11, space_after=2)
def note(t): para(t, size=8, space_after=8)

def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        v = kw.get(edge)
        if v:
            el = OxmlElement(f'w:{edge}'); el.set(qn('w:val'), v); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), '000000'); b.append(el)
    tcPr.append(b)

def table(rows, header_rows=1, col_widths=None, shade_first_col=True, font_size=9, bold_first_col=True):
    """rows: list of lists of str. header rows get bold + bottom border."""
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol); t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.style = 'Table Grid'
    for i, r in enumerate(rows):
        for j in range(ncol):
            v = r[j] if j < len(r) else ''
            c = t.cell(i, j); c.text = ''
            p = c.paragraphs[0]; run = p.add_run(str(v)); run.font.size = Pt(font_size); run.font.name = 'Arial'
            is_head = i < header_rows
            run.bold = is_head or (bold_first_col and j == 0 and str(v) != '')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (is_head or j > 0) else WD_ALIGN_PARAGRAPH.LEFT
            if is_head or (shade_first_col and j == 0):
                tcPr = c._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), 'EDEDED'); tcPr.append(sh)
    if col_widths:
        for j, w in enumerate(col_widths):
            for i in range(len(rows)): t.cell(i, j).width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def notes_table(syntax, n_rows=N, filt='<none>', cases='Statistics are based on all cases with valid data.', extra=None):
    rows = [['Notes', '', ''],
            ['Output Created', '', STAMP], ['Comments', '', ''],
            ['Input', 'Active Dataset', 'DataSet1'], ['', 'Filter', filt], ['', 'Weight', '<none>'], ['', 'Split File', '<none>'],
            ['', 'N of Rows in Working Data File', str(n_rows)],
            ['Missing Value Handling', 'Definition of Missing', 'User-defined missing values are treated as missing.'], ['', 'Cases Used', cases],
            ['Syntax', '', syntax],
            ['Resources', 'Processor Time', '00:00:00.02'], ['', 'Elapsed Time', '00:00:00.02']]
    if extra: rows += extra
    table(rows, header_rows=1, col_widths=[4.2, 5.0, 8.0], font_size=8)

SYNTAX = []  # collected for the .sps file
def syn(s): SYNTAX.append(s); return s

# =====================================================================
# Title / disclaimer
# =====================================================================
para('פלט SPSS לדוגמה: חשיפה למסכים, שינה ופעילות גופנית בקרב תלמידים עם לקויות למידה', bold=True, size=14, align=WD_ALIGN_PARAGRAPH.RIGHT)
para('נתונים מדומים לתרגול בלבד (N = 65). כל הטבלאות חושבו מקובץ הסימולציה simulated_data_N65_for_JASP.csv ואינן ממצאי מחקר. הפלט מדגים כיצד ייראה הניתוח עם הנתונים האמיתיים, ובכל בלוק מופיעה פקודת ה-SPSS (Syntax) להרצה על הקובץ האמיתי.', size=10, color=(192, 0, 0), align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=12)

# =====================================================================
# 1. FREQUENCIES (demographics)
# =====================================================================
proc_title('Frequencies')
notes_table(syn('FREQUENCIES VARIABLES=Group_LD Gender Grade Sector ADHD Accommodations\n  /ORDER=ANALYSIS.'))
para('[DataSet1]', size=9)
tbl_title('Statistics')
vars1 = ['Group_LD', 'Gender', 'Grade', 'Sector', 'ADHD', 'Accommodations']
table([['', ''] + [VLAB[v] for v in vars1], ['N', 'Valid'] + [str(N)] * len(vars1), ['', 'Missing'] + ['0'] * len(vars1)], header_rows=1, font_size=8)
tbl_title('Frequency Table')
for v in vars1:
    tbl_title(VLAB[v])
    vc = df[v].value_counts().sort_index(); cum = 0; rows = [['', '', 'Frequency', 'Percent', 'Valid Percent', 'Cumulative Percent']]
    for k, n in vc.items():
        cum += n; rows.append(['Valid' if k == vc.index[0] else '', VALLAB[v][k], str(n), f(100 * n / N, 1), f(100 * n / N, 1), f(100 * cum / N, 1)])
    rows.append(['', 'Total', str(N), '100.0', '100.0', ''])
    table(rows, col_widths=[1.5, 4, 2.2, 2.2, 2.4, 3.0])

# =====================================================================
# 2. DESCRIPTIVES
# =====================================================================
proc_title('Descriptives')
dvars = ['Screen_Leisure_WD', 'Screen_Leisure_WE', 'Screen_Leisure_Avg', 'Bedtime_WD', 'Sleep_Dur_WD', 'Sleep_Dur_WE', 'Weekend_Shift',
         'Sleepiness_Total', 'PA_Days60', 'PA_Hours_Vig', 'Grade_Avg', 'Grade_Math', 'Grade_English', 'Func_Total']
notes_table(syn('DESCRIPTIVES VARIABLES=' + ' '.join(dvars) + '\n  /STATISTICS=MEAN STDDEV MIN MAX SKEWNESS KURTOSIS.'), cases='All non-missing data are used.')
tbl_title('Descriptive Statistics')
n = N
se_skew = np.sqrt(6 * n * (n - 1) / ((n - 2) * (n + 1) * (n + 3))); se_kurt = 2 * se_skew * np.sqrt((n * n - 1) / ((n - 3) * (n + 5)))
rows = [['', 'N', 'Minimum', 'Maximum', 'Mean', 'Std. Deviation', 'Skewness', '', 'Kurtosis', ''],
        ['', 'Statistic', 'Statistic', 'Statistic', 'Statistic', 'Statistic', 'Statistic', 'Std. Error', 'Statistic', 'Std. Error']]
for v in dvars:
    s = df[v]
    rows.append([VLAB[v], str(n), f(s.min(), 2), f(s.max(), 2), f(s.mean(), 4), f(s.std(ddof=1), 5), f(s.skew(), 3), f(se_skew, 3), f(s.kurt(), 3), f(se_kurt, 3)])
rows.append(['Valid N (listwise)', str(n)] + [''] * 8)
table(rows, header_rows=2, col_widths=[5.2, 1.2, 1.6, 1.6, 1.7, 1.9, 1.5, 1.4, 1.5, 1.4], font_size=8)

# =====================================================================
# 3. RELIABILITY
# =====================================================================
def reliability_block(items, scale_name, syntax_items):
    proc_title('Reliability')
    notes_table(syn(f"RELIABILITY\n  /VARIABLES={syntax_items}\n  /SCALE('{scale_name}') ALL\n  /MODEL=ALPHA\n  /STATISTICS=DESCRIPTIVE CORR SCALE\n  /SUMMARY=TOTAL."),
                cases='Statistics are based on all cases with valid data for all variables in the procedure.', extra=[['', 'Matrix Input', '']])
    para(f'Scale: {scale_name}', bold=True, size=11)
    X = df[items].astype(float); k = len(items)
    tbl_title('Case Processing Summary')
    table([['', '', 'N', '%'], ['Cases', 'Valid', str(N), '100.0'], ['', 'Excludedᵃ', '0', '.0'], ['', 'Total', str(N), '100.0']], col_widths=[2, 2.5, 2, 2])
    note('a. Listwise deletion based on all variables in the procedure.')
    item_var = X.var(ddof=1); tot = X.sum(axis=1); alpha = k / (k - 1) * (1 - item_var.sum() / tot.var(ddof=1))
    R = X.corr(); rbar = (R.values.sum() - k) / (k * (k - 1)); alpha_std = k * rbar / (1 + (k - 1) * rbar)
    tbl_title('Reliability Statistics')
    table([["Cronbach's Alpha", "Cronbach's Alpha Based on Standardized Items", 'N of Items'], [f(alpha), f(alpha_std), str(k)]], col_widths=[3.5, 5.5, 2.5], shade_first_col=False, bold_first_col=False)
    tbl_title('Item Statistics')
    table([['', 'Mean', 'Std. Deviation', 'N']] + [[VLAB[i], f(X[i].mean(), 2), f(X[i].std(ddof=1), 3), str(N)] for i in items], col_widths=[6, 2, 2.6, 1.5])
    tbl_title('Inter-Item Correlation Matrix')
    table([[''] + [VLAB[i] for i in items]] + [[VLAB[i]] + [f(R.loc[i, j]) for j in items] for i in items], font_size=7)
    tbl_title('Item-Total Statistics')
    rows = [['', 'Scale Mean if Item Deleted', 'Scale Variance if Item Deleted', 'Corrected Item-Total Correlation', 'Squared Multiple Correlation', "Cronbach's Alpha if Item Deleted"]]
    for i in items:
        rest = [j for j in items if j != i]; tr = X[rest].sum(axis=1)
        citc = np.corrcoef(X[i], tr)[0, 1]
        Xo = sm.add_constant(X[rest]); smc = sm.OLS(X[i], Xo).fit().rsquared
        a_del = (k - 1) / (k - 2) * (1 - X[rest].var(ddof=1).sum() / tr.var(ddof=1))
        rows.append([VLAB[i], f(tr.mean(), 2), f(tr.var(ddof=1), 3), f(citc), f(smc), f(a_del)])
    table(rows, col_widths=[5, 2, 2.2, 2.2, 2.2, 2.4], font_size=8)
    tbl_title('Scale Statistics')
    table([['Mean', 'Variance', 'Std. Deviation', 'N of Items'], [f(tot.mean(), 2), f(tot.var(ddof=1), 3), f(tot.std(ddof=1), 3), str(k)]], shade_first_col=False, bold_first_col=False, col_widths=[2.5, 2.5, 2.8, 2.2])
    return alpha

a1 = reliability_block(['S1', 'S2', 'S3R', 'S4', 'S5', 'S6', 'S7', 'S8'], 'ישנוניות יומית', 'S1 S2 S3R S4 S5 S6 S7 S8')
a2 = reliability_block(['F1', 'F2', 'F3', 'F4R', 'F5'], 'תפקוד יומי בלמידה', 'F1 F2 F3 F4R F5')

# =====================================================================
# 4. EXPLORE: normality
# =====================================================================
proc_title('Explore')
evars = ['Screen_Leisure_Avg', 'Sleep_Dur_WD', 'Weekend_Shift', 'Sleepiness_Total', 'PA_Days60', 'Grade_Avg', 'Func_Total']
notes_table(syn('EXAMINE VARIABLES=' + ' '.join(evars) + '\n  /PLOT HISTOGRAM NPPLOT\n  /STATISTICS DESCRIPTIVES\n  /MISSING PAIRWISE.'),
            cases='Statistics are based on cases with no missing values for the dependent variable or factor(s) being analyzed.')
para('Total Sample', bold=True, size=11)
tbl_title('Tests of Normality')
rows = [['', 'Kolmogorov-Smirnovᵃ', '', '', 'Shapiro-Wilk', '', ''], ['', 'Statistic', 'df', 'Sig.', 'Statistic', 'df', 'Sig.']]
for v in evars:
    ks, ksp = lilliefors(df[v], dist='norm'); w, wp = stats.shapiro(df[v])
    rows.append([VLAB[v], f(ks), str(N), fp(ksp), f(w), str(N), fp(wp)])
table(rows, header_rows=2, col_widths=[5.5, 1.8, 1.2, 1.5, 1.8, 1.2, 1.5])
note('a. Lilliefors Significance Correction')

# =====================================================================
# 5. CROSSTABS: guidelines by group
# =====================================================================
proc_title('Crosstabs')
notes_table(syn('CROSSTABS\n  /TABLES=Group_LD BY Meet_Screen Meet_Sleep Meet_PA Guidelines_Grp\n  /FORMAT=AVALUE TABLES\n  /STATISTICS=CHISQ PHI\n  /CELLS=COUNT ROW\n  /COUNT ROUND CELL.'),
            cases='Statistics for each table are based on all the cases with valid data in the specified range(s) for all variables in each table.')
tbl_title('Case Processing Summary')
rows = [['', 'Cases', '', '', '', '', ''], ['', 'Valid', '', 'Missing', '', 'Total', ''], ['', 'N', 'Percent', 'N', 'Percent', 'N', 'Percent']]
for v in ['Meet_Screen', 'Meet_Sleep', 'Meet_PA', 'Guidelines_Grp']:
    rows.append([f"{VLAB['Group_LD']} * {VLAB[v]}", str(N), '100.0%', '0', '0.0%', str(N), '100.0%'])
table(rows, header_rows=3, col_widths=[6.5, 1.3, 1.6, 1.3, 1.6, 1.3, 1.6], font_size=8)
for v in ['Meet_Screen', 'Meet_Sleep', 'Meet_PA', 'Guidelines_Grp']:
    ct = pd.crosstab(df['Group_LD'], df[v]); cols = list(ct.columns)
    tbl_title(f"{VLAB['Group_LD']} * {VLAB[v]} Crosstabulation")
    rows = [['', '', '', VLAB[v]] + [''] * (len(cols) - 1) + ['Total'], ['', '', ''] + [VALLAB[v][c] for c in cols] + ['']]
    for g in [0, 1]:
        rt = ct.loc[g].sum()
        rows.append([VLAB['Group_LD'] if g == 0 else '', VALLAB['Group_LD'][g], 'Count'] + [str(ct.loc[g, c]) for c in cols] + [str(rt)])
        rows.append(['', '', f"% within {VLAB['Group_LD']}"] + [f(100 * ct.loc[g, c] / rt, 1) + '%' for c in cols] + ['100.0%'])
    rows.append(['Total', '', 'Count'] + [str(ct[c].sum()) for c in cols] + [str(N)])
    rows.append(['', '', f"% within {VLAB['Group_LD']}"] + [f(100 * ct[c].sum() / N, 1) + '%' for c in cols] + ['100.0%'])
    table(rows, header_rows=2, font_size=8)
    chi2, p, dof, exp = stats.chi2_contingency(ct.values, correction=False)
    g2, gp, _, _ = stats.chi2_contingency(ct.values, correction=False, lambda_='log-likelihood')
    tbl_title('Chi-Square Tests')
    rows = [['', 'Value', 'df', 'Asymptotic Significance (2-sided)', 'Exact Sig. (2-sided)', 'Exact Sig. (1-sided)'],
            ['Pearson Chi-Square', f(chi2) + 'ᵃ', str(dof), fp(p), '', '']]
    if ct.shape == (2, 2):
        cc, cp, _, _ = stats.chi2_contingency(ct.values, correction=True)
        _, fe = stats.fisher_exact(ct.values); fe1 = min(stats.fisher_exact(ct.values, alternative='less')[1], stats.fisher_exact(ct.values, alternative='greater')[1])
        rows.append(['Continuity Correctionᵇ', f(cc), '1', fp(cp), '', ''])
        rows.append(['Likelihood Ratio', f(g2), str(dof), fp(gp), '', ''])
        rows.append(["Fisher's Exact Test", '', '', '', fp(fe), fp(fe1)])
    else:
        rows.append(['Likelihood Ratio', f(g2), str(dof), fp(gp), '', ''])
    rows.append(['N of Valid Cases', str(N), '', '', '', ''])
    table(rows, col_widths=[3.8, 1.8, 1.0, 3.0, 2.4, 2.4], font_size=8)
    small = (exp < 5).sum(); minexp = exp.min()
    note(f"a. {small} cells ({100 * small / exp.size:.1f}%) have expected count less than 5. The minimum expected count is {minexp:.2f}." + ("\nb. Computed only for a 2x2 table" if ct.shape == (2, 2) else ''))
    phi = np.sqrt(chi2 / N)
    tbl_title('Symmetric Measures')
    table([['', '', 'Value', 'Approximate Significance'], ['Nominal by Nominal', 'Phi' if ct.shape == (2, 2) else "Cramer's V", f(phi if ct.shape == (2, 2) else np.sqrt(chi2 / (N * (min(ct.shape) - 1)))), fp(p)], ['N of Valid Cases', '', str(N), '']], col_widths=[3.5, 2.5, 2, 3.5])

# =====================================================================
# 6. CORRELATIONS (LD group)
# =====================================================================
ld = df[df.Group_LD == 1]; nld = len(ld)
cvars = ['Screen_Leisure_Avg', 'Bedtime_WD', 'Sleep_Dur_WD', 'Weekend_Shift', 'Sleepiness_Total', 'PA_Days60', 'Grade_Avg', 'Func_Total']
def corr_block(method):
    proc_title('Correlations' if method == 'pearson' else 'Nonparametric Correlations')
    if method == 'pearson':
        s = syn('USE ALL.\nCOMPUTE filter_LD=(Group_LD = 1).\nFILTER BY filter_LD.\nEXECUTE.\nCORRELATIONS\n  /VARIABLES=' + ' '.join(cvars) + '\n  /PRINT=TWOTAIL NOSIG\n  /STATISTICS DESCRIPTIVES\n  /MISSING=PAIRWISE.')
    else:
        s = syn('NONPAR CORR\n  /VARIABLES=' + ' '.join(cvars) + '\n  /PRINT=SPEARMAN TWOTAIL NOSIG\n  /MISSING=PAIRWISE.')
    notes_table(s, n_rows=nld, filt='filter_LD', cases='Statistics for each pair of variables are based on all the cases with valid data for that pair.')
    para('[DataSet1] קבוצת המחקר בלבד: תלמידים עם לקות למידה (Group_LD = 1)', size=9)
    if method == 'pearson':
        tbl_title('Descriptive Statistics')
        table([['', 'Mean', 'Std. Deviation', 'N']] + [[VLAB[v], f(ld[v].mean(), 4), f(ld[v].std(ddof=1), 5), str(nld)] for v in cvars], col_widths=[6, 2, 2.6, 1.5])
    tbl_title('Correlations')
    head = [['', ''] + [VLAB[v] for v in cvars]]; rows = []
    for a in cvars:
        r1, r2, r3 = [VLAB[a], 'Pearson Correlation' if method == 'pearson' else 'Correlation Coefficient'], ['', 'Sig. (2-tailed)'], ['', 'N']
        for b in cvars:
            if a == b: r1.append('1.000' if method != 'pearson' else '1'); r2.append(''); r3.append(str(nld)); continue
            r, p = (stats.pearsonr(ld[a], ld[b]) if method == 'pearson' else stats.spearmanr(ld[a], ld[b]))
            r1.append(f(r) + stars(p)); r2.append(fp(p)); r3.append(str(nld))
        rows += [r1, r2, r3]
    if method == 'spearman':
        rows = [["Spearman's rho"] + r[1:] if i == 0 else r for i, r in enumerate(rows)]
    table(head + rows, header_rows=1, font_size=7)
    note('**. Correlation is significant at the 0.01 level (2-tailed).\n*. Correlation is significant at the 0.05 level (2-tailed).')
corr_block('pearson'); corr_block('spearman')
syn('FILTER OFF.\nUSE ALL.\nEXECUTE.')

# =====================================================================
# 7. T-TEST by group
# =====================================================================
proc_title('T-Test')
tvars = ['Screen_Leisure_Avg', 'Sleep_Dur_WD', 'Weekend_Shift', 'Sleepiness_Total', 'PA_Days60', 'Grade_Avg', 'Grade_Math', 'Grade_English', 'Func_Total']
notes_table(syn('T-TEST GROUPS=Group_LD(0 1)\n  /MISSING=ANALYSIS\n  /VARIABLES=' + ' '.join(tvars) + '\n  /ES DISPLAY(TRUE)\n  /CRITERIA=CI(.95).'),
            cases='Statistics for each analysis are based on the cases with no missing or out-of-range data for any variable in the analysis.')
tbl_title('Group Statistics')
rows = [['', VLAB['Group_LD'], 'N', 'Mean', 'Std. Deviation', 'Std. Error Mean']]
g0 = df[df.Group_LD == 0]; g1 = df[df.Group_LD == 1]
for v in tvars:
    for k, g in ((0, g0), (1, g1)):
        rows.append([VLAB[v] if k == 0 else '', VALLAB['Group_LD'][k], str(len(g)), f(g[v].mean(), 4), f(g[v].std(ddof=1), 5), f(g[v].sem(), 5)])
table(rows, col_widths=[5, 3, 1.2, 2, 2.4, 2.4], font_size=8)
tbl_title('Independent Samples Test')
rows = [['', '', "Levene's Test for Equality of Variances", '', 't-test for Equality of Means', '', '', '', '', '', '', ''],
        ['', '', 'F', 'Sig.', 't', 'df', 'Sig. (2-tailed)', 'Mean Difference', 'Std. Error Difference', '95% CI Lower', '95% CI Upper', '']]
es_rows = [['', '', 'Standardizerᵃ', 'Point Estimate', '95% CI Lower', '95% CI Upper']]
for v in tvars:
    a, b = g0[v].astype(float), g1[v].astype(float)
    lf, lp = stats.levene(a, b, center='mean')
    t1 = stats.ttest_ind(a, b, equal_var=True); t2 = stats.ttest_ind(a, b, equal_var=False)
    md = a.mean() - b.mean(); n1, n2 = len(a), len(b)
    sp = np.sqrt(((n1 - 1) * a.var(ddof=1) + (n2 - 1) * b.var(ddof=1)) / (n1 + n2 - 2)); se1 = sp * np.sqrt(1 / n1 + 1 / n2)
    se2 = np.sqrt(a.var(ddof=1) / n1 + b.var(ddof=1) / n2); df2 = t2.df
    ci1 = stats.t.ppf(.975, n1 + n2 - 2) * se1; ci2 = stats.t.ppf(.975, df2) * se2
    rows.append([VLAB[v], 'Equal variances assumed', f(lf), fp(lp), f(t1.statistic), str(n1 + n2 - 2), fp(t1.pvalue), f(md, 4), f(se1, 4), f(md - ci1, 4), f(md + ci1, 4), ''])
    rows.append(['', 'Equal variances not assumed', '', '', f(t2.statistic), f(df2), fp(t2.pvalue), f(md, 4), f(se2, 4), f(md - ci2, 4), f(md + ci2, 4), ''])
    d = md / sp; J = 1 - 3 / (4 * (n1 + n2) - 9); se_d = np.sqrt((n1 + n2) / (n1 * n2) + d * d / (2 * (n1 + n2)))
    es_rows.append([VLAB[v], "Cohen's d", f(sp, 4), f(d), f(d - 1.96 * se_d), f(d + 1.96 * se_d)])
    es_rows.append(['', "Hedges' correction", f(sp / J, 4), f(d * J), f((d - 1.96 * se_d) * J), f((d + 1.96 * se_d) * J)])
table(rows, header_rows=2, font_size=7)
tbl_title('Independent Samples Effect Sizes')
table(es_rows, col_widths=[4.5, 3, 2, 2, 2, 2], font_size=8)
note("a. The denominator used in estimating the effect sizes. Cohen's d uses the pooled standard deviation. Hedges' correction uses the pooled standard deviation, plus a correction factor.")

# =====================================================================
# 8. ONEWAY: outcome by number of guidelines
# =====================================================================
proc_title('Oneway')
notes_table(syn('ONEWAY Grade_Avg Func_Total BY Guidelines_Grp\n  /STATISTICS DESCRIPTIVES HOMOGENEITY WELCH\n  /MISSING ANALYSIS\n  /POSTHOC=BONFERRONI ALPHA(0.05).'),
            cases='Statistics for each analysis are based on cases with no missing data for any variable in the analysis.')
groups = [0, 1, 2]
tbl_title('Descriptives')
rows = [['', '', 'N', 'Mean', 'Std. Deviation', 'Std. Error', '95% CI Lower Bound', '95% CI Upper Bound', 'Minimum', 'Maximum']]
for v in ['Grade_Avg', 'Func_Total']:
    for g in groups:
        s = df[df.Guidelines_Grp == g][v]; h = stats.t.ppf(.975, len(s) - 1) * s.sem()
        rows.append([VLAB[v] if g == 0 else '', VALLAB['Guidelines_Grp'][g], str(len(s)), f(s.mean(), 4), f(s.std(ddof=1), 5), f(s.sem(), 5), f(s.mean() - h, 4), f(s.mean() + h, 4), f(s.min(), 2), f(s.max(), 2)])
    s = df[v]; h = stats.t.ppf(.975, N - 1) * s.sem()
    rows.append(['', 'Total', str(N), f(s.mean(), 4), f(s.std(ddof=1), 5), f(s.sem(), 5), f(s.mean() - h, 4), f(s.mean() + h, 4), f(s.min(), 2), f(s.max(), 2)])
table(rows, font_size=7)
tbl_title('Test of Homogeneity of Variances')
rows = [['', '', 'Levene Statistic', 'df1', 'df2', 'Sig.']]
for v in ['Grade_Avg', 'Func_Total']:
    samples = [df[df.Guidelines_Grp == g][v] for g in groups]
    lf, lp = stats.levene(*samples, center='mean'); rows.append([VLAB[v], 'Based on Mean', f(lf), '2', str(N - 3), fp(lp)])
    lf, lp = stats.levene(*samples, center='median'); rows.append(['', 'Based on Median', f(lf), '2', str(N - 3), fp(lp)])
table(rows, col_widths=[5, 3, 2.2, 1.2, 1.2, 1.5], font_size=8)
tbl_title('ANOVA')
rows = [['', '', 'Sum of Squares', 'df', 'Mean Square', 'F', 'Sig.']]
posthoc = []
for v in ['Grade_Avg', 'Func_Total']:
    samples = [df[df.Guidelines_Grp == g][v].astype(float) for g in groups]
    grand = df[v].mean(); ssb = sum(len(s) * (s.mean() - grand) ** 2 for s in samples); ssw = sum(((s - s.mean()) ** 2).sum() for s in samples)
    dfb, dfw = 2, N - 3; msb, msw = ssb / dfb, ssw / dfw; F = msb / msw; p = 1 - stats.f.cdf(F, dfb, dfw)
    rows += [[VLAB[v], 'Between Groups', f(ssb, 3), '2', f(msb, 3), f(F), fp(p)], ['', 'Within Groups', f(ssw, 3), str(dfw), f(msw, 3), '', ''], ['', 'Total', f(ssb + ssw, 3), str(N - 1), '', '', '']]
    for i in range(3):
        for j in range(3):
            if i == j: continue
            si, sj = samples[i], samples[j]; mdiff = si.mean() - sj.mean(); se = np.sqrt(msw * (1 / len(si) + 1 / len(sj)))
            t = mdiff / se; pb = min(1.0, 3 * 2 * (1 - stats.t.cdf(abs(t), dfw))); crit = stats.t.ppf(1 - .05 / (2 * 3), dfw)
            posthoc.append([VLAB[v] if (i == 0 and j == 1) else '', VALLAB['Guidelines_Grp'][i] if j == (0 if i else 1) else '', VALLAB['Guidelines_Grp'][j], f(mdiff, 4) + ('*' if pb < .05 else ''), f(se, 4), fp(pb), f(mdiff - crit * se, 4), f(mdiff + crit * se, 4)])
table(rows, col_widths=[4.5, 2.8, 2.4, 1.2, 2.2, 1.6, 1.5], font_size=8)
tbl_title('Robust Tests of Equality of Means')
rows = [['', '', 'Statisticᵃ', 'df1', 'df2', 'Sig.']]
for v in ['Grade_Avg', 'Func_Total']:
    samples = [df[df.Guidelines_Grp == g][v].astype(float) for g in groups]
    w = [len(s) / s.var(ddof=1) for s in samples]; sw = sum(w); mw = sum(wi * s.mean() for wi, s in zip(w, samples)) / sw
    k = 3; num = sum(wi * (s.mean() - mw) ** 2 for wi, s in zip(w, samples)) / (k - 1)
    lam = 3 * sum((1 - wi / sw) ** 2 / (len(s) - 1) for wi, s in zip(w, samples)) / (k * k - 1); Fw = num / (1 + 2 * (k - 2) * lam / 3); df2w = 1 / lam
    rows.append([VLAB[v], 'Welch', f(Fw), '2', f(df2w), fp(1 - stats.f.cdf(Fw, 2, df2w))])
table(rows, col_widths=[5, 2, 2, 1.2, 2, 1.5], font_size=8)
note('a. Asymptotically F distributed.')
para('Post Hoc Tests', bold=True, size=11)
tbl_title('Multiple Comparisons')
para('Bonferroni', size=9)
table([['Dependent Variable', '(I) ' + VLAB['Guidelines_Grp'], '(J) ' + VLAB['Guidelines_Grp'], 'Mean Difference (I-J)', 'Std. Error', 'Sig.', '95% CI Lower Bound', '95% CI Upper Bound']] + posthoc, font_size=7)
note('*. The mean difference is significant at the 0.05 level.')

# =====================================================================
# 9. REGRESSION (hierarchical)
# =====================================================================
proc_title('Regression')
preds1 = ['Screen_Leisure_Avg', 'Sleep_Dur_WD', 'Weekend_Shift', 'PA_Days60', 'Sleepiness_Total']; preds2 = ['Group_LD', 'Gender']
notes_table(syn('REGRESSION\n  /DESCRIPTIVES MEAN STDDEV CORR SIG N\n  /MISSING LISTWISE\n  /STATISTICS COEFF OUTS CI(95) R ANOVA CHANGE COLLIN TOL\n  /CRITERIA=PIN(.05) POUT(.10)\n  /NOORIGIN\n  /DEPENDENT Grade_Avg\n  /METHOD=ENTER ' + ' '.join(preds1) + '\n  /METHOD=ENTER ' + ' '.join(preds2) + '\n  /SCATTERPLOT=(*ZRESID ,*ZPRED)\n  /RESIDUALS HISTOGRAM(ZRESID) NORMPROB(ZRESID).'),
            cases='Statistics are based on cases with no missing values for any variable used.')
y = df['Grade_Avg'].astype(float)
X1 = sm.add_constant(df[preds1].astype(float)); X2 = sm.add_constant(df[preds1 + preds2].astype(float))
m1 = sm.OLS(y, X1).fit(); m2 = sm.OLS(y, X2).fit()
tbl_title('Variables Entered/Removedᵃ')
table([['Model', 'Variables Entered', 'Variables Removed', 'Method'], ['1', ', '.join(VLAB[p] for p in preds1) + 'ᵇ', '.', 'Enter'], ['2', ', '.join(VLAB[p] for p in preds2) + 'ᵇ', '.', 'Enter']], col_widths=[1.2, 9, 2.5, 1.8], font_size=8)
note(f"a. Dependent Variable: {VLAB['Grade_Avg']}\nb. All requested variables entered.")
tbl_title('Model Summaryᶜ')
r2c = m2.rsquared - m1.rsquared; df1c = len(preds2); df2c = int(m2.df_resid); Fc = (r2c / df1c) / ((1 - m2.rsquared) / df2c); pc = 1 - stats.f.cdf(Fc, df1c, df2c)
table([['Model', 'R', 'R Square', 'Adjusted R Square', 'Std. Error of the Estimate', 'R Square Change', 'F Change', 'df1', 'df2', 'Sig. F Change'],
       ['1', f(np.sqrt(m1.rsquared)) + 'ᵃ', f(m1.rsquared), f(m1.rsquared_adj), f(np.sqrt(m1.mse_resid), 4), f(m1.rsquared), f(m1.fvalue), str(len(preds1)), str(int(m1.df_resid)), fp(m1.f_pvalue)],
       ['2', f(np.sqrt(m2.rsquared)) + 'ᵇ', f(m2.rsquared), f(m2.rsquared_adj), f(np.sqrt(m2.mse_resid), 4), f(r2c), f(Fc), str(df1c), str(df2c), fp(pc)]], font_size=7)
note(f"a. Predictors: (Constant), {', '.join(VLAB[p] for p in preds1)}\nb. Predictors: (Constant), {', '.join(VLAB[p] for p in preds1 + preds2)}\nc. Dependent Variable: {VLAB['Grade_Avg']}")
tbl_title('ANOVAᵃ')
rows = [['Model', '', 'Sum of Squares', 'df', 'Mean Square', 'F', 'Sig.']]
for k, m in ((1, m1), (2, m2)):
    rows += [[str(k), 'Regression', f(m.ess, 3), str(int(m.df_model)), f(m.ess / m.df_model, 3), f(m.fvalue), fp(m.f_pvalue) + ('ᵇ' if k == 1 else 'ᶜ')], ['', 'Residual', f(m.ssr, 3), str(int(m.df_resid)), f(m.mse_resid, 3), '', ''], ['', 'Total', f(m.ess + m.ssr, 3), str(N - 1), '', '', '']]
table(rows, col_widths=[1.2, 2.5, 2.6, 1.2, 2.4, 1.6, 1.5], font_size=8)
note(f"a. Dependent Variable: {VLAB['Grade_Avg']}\nb. Predictors: (Constant), {', '.join(VLAB[p] for p in preds1)}\nc. Predictors: (Constant), {', '.join(VLAB[p] for p in preds1 + preds2)}")
tbl_title('Coefficientsᵃ')
rows = [['Model', '', 'Unstandardized B', 'Std. Error', 'Standardized Beta', 't', 'Sig.', '95% CI Lower', '95% CI Upper', 'Tolerance', 'VIF']]
for k, m, preds in ((1, m1, preds1), (2, m2, preds1 + preds2)):
    ci = m.conf_int(); Xp = df[preds].astype(float)
    for i, name in enumerate(['const'] + preds):
        beta = '' if name == 'const' else f(m.params[name] * Xp[name].std(ddof=1) / y.std(ddof=1))
        if name == 'const': tol = vif = ''
        else:
            others = [p for p in preds if p != name]; r2 = sm.OLS(Xp[name], sm.add_constant(Xp[others])).fit().rsquared; tol = f(1 - r2); vif = f(1 / (1 - r2))
        rows.append([str(k) if i == 0 else '', '(Constant)' if name == 'const' else VLAB[name], f(m.params[name], 3), f(m.bse[name], 3), beta, f(m.tvalues[name]), fp(m.pvalues[name]), f(ci.loc[name, 0], 3), f(ci.loc[name, 1], 3), tol, vif])
table(rows, font_size=7)
note(f"a. Dependent Variable: {VLAB['Grade_Avg']}")
tbl_title('Residuals Statisticsᵃ')
res = m2.resid; pred = m2.fittedvalues; zres = res / np.sqrt(m2.mse_resid); zpred = (pred - pred.mean()) / pred.std(ddof=1)
table([['', 'Minimum', 'Maximum', 'Mean', 'Std. Deviation', 'N'],
       ['Predicted Value', f(pred.min(), 4), f(pred.max(), 4), f(pred.mean(), 4), f(pred.std(ddof=1), 5), str(N)],
       ['Residual', f(res.min(), 5), f(res.max(), 5), '.00000', f(res.std(ddof=1), 5), str(N)],
       ['Std. Predicted Value', f(zpred.min()), f(zpred.max()), '.000', '1.000', str(N)],
       ['Std. Residual', f(zres.min()), f(zres.max()), '.000', f(zres.std(ddof=1)), str(N)]], col_widths=[4, 2, 2, 2, 2.5, 1.2], font_size=8)
note(f"a. Dependent Variable: {VLAB['Grade_Avg']}")

# =====================================================================
# Summary paragraph in Hebrew (how to read)
# =====================================================================
para('הערה לקריאת הפלט', bold=True, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT)
para('כל בלוק פותח בטבלת Notes עם פקודת ה-Syntax, כמו בפלט SPSS מקורי. כדי לחזור על הניתוח עם הנתונים האמיתיים: לפתוח את קובץ ה-SAV המצורף (או להקליד את הנתונים באותו מבנה), לפתוח את קובץ ה-SPS ב-File > Open > Syntax, ולהריץ Run > All. הפלט שיתקבל יהיה זהה במבנהו לפלט זה, עם המספרים האמיתיים. הנתונים כאן מדומים ואינם ממצאים.', size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)

doc.save('פלט_SPSS_לדוגמה_נתונים_מדומים.docx')

# ---------- .sav with labels ----------
sav_cols = [c for c in df.columns if c not in ('S3R', 'F4R', 'Guidelines_Grp')] + ['S3R', 'F4R', 'Guidelines_Grp']
pyreadstat.write_sav(df[sav_cols], 'simulated_data_N65.sav', column_labels={c: VLAB.get(c, c) for c in sav_cols},
                     variable_value_labels={k: {float(a): b for a, b in v.items()} for k, v in VALLAB.items() if k in sav_cols})

# ---------- .sps syntax ----------
hdr = ['* Syntax for the seminar study: screens, sleep and physical activity among students with learning disabilities.',
       '* Run on the real data file (same column structure as the simulated file).', '',
       'GET FILE=\'simulated_data_N65.sav\'.', 'DATASET NAME DataSet1 WINDOW=FRONT.', '',
       '* Reverse-coded items and grouped guideline variable.',
       'COMPUTE S3R=4 - S3.', 'COMPUTE F4R=6 - F4.', 'COMPUTE Guidelines_Grp=Guidelines_Count.', 'RECODE Guidelines_Grp (3=2).',
       'VARIABLE LABELS S3R \'ערני רוב היום (מהופך)\' F4R \'מרגיש עייף בשיעורים (מהופך)\' Guidelines_Grp \'מספר המלצות (מקובץ)\'.',
       'VALUE LABELS Guidelines_Grp 0 \'אף המלצה\' 1 \'המלצה אחת\' 2 \'שתיים או שלוש\'.', 'EXECUTE.', '']
open('seminar_analysis_syntax.sps', 'w', encoding='utf-8').write('\n'.join(hdr) + '\n\n'.join(SYNTAX) + '\n')
print('done; alpha sleepiness', round(a1, 3), 'alpha function', round(a2, 3))
