# -*- coding: utf-8 -*-
"""Builds a SIMULATED dataset (N=65) matching the seminar questionnaire, for practicing JASP analyses."""
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

rng = np.random.default_rng(2026)
N = 65
n_ld = 40
rows = []

def clip(x, lo, hi): return float(min(max(x, lo), hi))

for i in range(N):
    ld = 1 if i < n_ld else 0
    gender = int(rng.choice([1, 2], p=[0.55, 0.45] if ld else [0.48, 0.52]))
    grade = int(rng.choice([7, 8, 9, 10, 11, 12]))
    age = grade + 5 + int(rng.choice([0, 1], p=[0.7, 0.3]))
    sector = int(rng.choice([1, 2, 3], p=[0.62, 0.32, 0.06]))
    if ld:
        ld_read = int(rng.random() < 0.7); ld_write = int(rng.random() < 0.5); ld_math = int(rng.random() < 0.4)
        if ld_read + ld_write + ld_math == 0: ld_read = 1
        accom = int(rng.random() < 0.85)
        adhd = int(rng.random() < 0.45)
    else:
        ld_read = ld_write = ld_math = 0; accom = 0
        adhd = int(rng.random() < 0.08)
    meds = int(adhd and rng.random() < 0.55)

    # latent screen propensity (hours/day leisure), higher in LD and in Arab sector, per literature
    base_screen = 3.4 + 0.4 * ld + 0.5 * (sector == 2) + 0.12 * (grade - 9) + rng.normal(0, 1.4)
    base_screen = clip(base_screen, 0.5, 9.0)
    tv_wd = clip(rng.normal(base_screen * 0.35, 0.5), 0, 6)
    games_wd = clip(rng.normal(base_screen * (0.28 if gender == 1 else 0.12), 0.5), 0, 6)
    social_wd = clip(rng.normal(base_screen * (0.3 if gender == 1 else 0.45), 0.5), 0, 7)
    study_wd = clip(rng.normal(1.2 - 0.2 * ld, 0.5), 0, 4)
    we_factor = rng.normal(1.45, 0.2)
    tv_we = clip(tv_wd * we_factor, 0, 8); games_we = clip(games_wd * we_factor, 0, 8)
    social_we = clip(social_wd * we_factor, 0, 9); study_we = clip(study_wd * 0.6, 0, 4)
    r = lambda v: round(v * 2) / 2  # half-hour resolution
    tv_wd, games_wd, social_wd, study_wd = map(r, (tv_wd, games_wd, social_wd, study_wd))
    tv_we, games_we, social_we, study_we = map(r, (tv_we, games_we, social_we, study_we))
    leis_wd = tv_wd + games_wd + social_wd
    leis_we = tv_we + games_we + social_we
    leis_avg = round((5 * leis_wd + 2 * leis_we) / 7, 2)

    bedroom = int(rng.random() < clip(0.35 + 0.06 * leis_avg, 0, 0.95))
    before_bed = int(clip(round(rng.normal(1.5 + 0.35 * leis_avg, 0.9)), 0, 4))
    night_wake = int(clip(round(rng.normal(0.3 + 0.25 * leis_avg + 0.5 * bedroom, 0.8)), 0, 3))

    # sleep: later bedtime with more screens and with LD/ADHD
    bed_wd = clip(rng.normal(22.1 + 0.14 * leis_avg + 0.25 * ld + 0.2 * adhd + 0.1 * (grade - 9), 0.6), 21.0, 26.0)
    wake_wd = clip(rng.normal(7.0, 0.35), 6.0, 8.0)
    bed_we = clip(bed_wd + rng.normal(1.3 + 0.1 * leis_avg, 0.6), 21.5, 28.0)
    wake_we = clip(wake_wd + rng.normal(2.5, 0.9), 7.0, 13.0)
    q = lambda v: round(v * 4) / 4  # 15-minute resolution
    bed_wd, wake_wd, bed_we, wake_we = map(q, (bed_wd, wake_wd, bed_we, wake_we))
    sleep_wd = round(wake_wd + 24 - bed_wd, 2)
    sleep_we = round(wake_we + 24 - bed_we, 2)
    shift = round(bed_we - bed_wd, 2)
    onset = int(clip(round(rng.normal(1.8 + 0.5 * ld + 0.3 * adhd + 0.15 * before_bed, 0.8)), 1, 4))

    # sleepiness items 0-4 (S3 reversed: high = alert)
    lat_sleepy = 1.3 + 0.4 * ld + 0.3 * adhd - 0.3 * (sleep_wd - 8.0) + 0.1 * leis_avg + rng.normal(0, 0.45)
    S = [int(clip(round(lat_sleepy + rng.normal(0, 0.7)), 0, 4)) for _ in range(8)]
    S[2] = int(clip(round(4 - lat_sleepy + rng.normal(0, 0.7)), 0, 4))  # reversed item
    sleepy_total = sum(S[:2]) + (4 - S[2]) + sum(S[3:])

    # physical activity
    lat_pa = 3.9 - 0.6 * ld - 0.18 * leis_avg + 0.8 * (gender == 1) + rng.normal(0, 1.7)
    pa_days = int(clip(round(lat_pa), 0, 7))
    pa_freq = int(clip(round(7 - (lat_pa + rng.normal(0, 0.8)) * 0.8), 1, 7))  # 1=every day ... 7=never
    pa_hours = int(clip(round(1 + (lat_pa + rng.normal(0, 0.8)) * 0.7), 1, 6))
    pa_club = int(rng.choice([0, 1, 2], p=[0.55, 0.3, 0.15] if ld else [0.4, 0.35, 0.25]))
    pe_active = int(clip(round(2.3 + 0.3 * lat_pa / 3 - 0.3 * ld + rng.normal(0, 0.9)), 0, 4))
    pe_like = int(clip(round(3.2 + 0.25 * lat_pa / 3 - 0.3 * ld + rng.normal(0, 0.9)), 1, 5))
    pa_sick = int(rng.random() < 0.1)

    # academic functioning
    lat_grade = 84 - 5 * ld - 1.2 * leis_avg - 1.0 * shift + 0.6 * pa_days - 0.4 * (sleepy_total - 12) + rng.normal(0, 6)
    g_avg = int(clip(round(lat_grade), 50, 100))
    g_math = int(clip(round(lat_grade - 3 - 4 * ld_math - 1.0 * leis_avg + rng.normal(0, 8)), 40, 100))
    g_eng = int(clip(round(lat_grade - 1 - 3 * ld_read - 0.5 * leis_avg + rng.normal(0, 8)), 40, 100))
    lat_func = 3.7 - 0.35 * ld - 0.08 * leis_avg - 0.05 * (sleepy_total - 12) + 0.05 * pa_days + rng.normal(0, 0.4)
    F = [int(clip(round(lat_func + rng.normal(0, 0.7)), 1, 5)) for _ in range(5)]
    F[3] = int(clip(round(6 - lat_func + rng.normal(0, 0.7)), 1, 5))  # reversed: tired in class
    func_total = round((F[0] + F[1] + F[2] + (6 - F[3]) + F[4]) / 5, 2)

    meet_screen = int(leis_avg <= 2)
    sleep_lo, sleep_hi = (9, 11) if age <= 13 else (8, 10)
    meet_sleep = int(sleep_lo <= sleep_wd <= sleep_hi)
    meet_pa = int(pa_days == 7)
    guidelines = meet_screen + meet_sleep + meet_pa

    rows.append(dict(
        ID=i + 1, Group_LD=ld, Gender=gender, Age=age, Grade=grade, Sector=sector,
        LD_Reading=ld_read, LD_Writing=ld_write, LD_Math=ld_math, Accommodations=accom, ADHD=adhd, Medication=meds,
        TV_WD=tv_wd, TV_WE=tv_we, Games_WD=games_wd, Games_WE=games_we, Social_WD=social_wd, Social_WE=social_we,
        Study_WD=study_wd, Study_WE=study_we, Screen_Leisure_WD=leis_wd, Screen_Leisure_WE=leis_we, Screen_Leisure_Avg=leis_avg,
        Screen_Bedroom=bedroom, Screen_BeforeBed=before_bed, Night_Wake=night_wake,
        Bedtime_WD=bed_wd, Waketime_WD=wake_wd, Bedtime_WE=bed_we, Waketime_WE=wake_we,
        Sleep_Dur_WD=sleep_wd, Sleep_Dur_WE=sleep_we, Weekend_Shift=shift, Sleep_Onset=onset,
        S1=S[0], S2=S[1], S3=S[2], S4=S[3], S5=S[4], S6=S[5], S7=S[6], S8=S[7], Sleepiness_Total=sleepy_total,
        PA_Days60=pa_days, PA_Freq_Vig=pa_freq, PA_Hours_Vig=pa_hours, PA_Club=pa_club, PE_Active=pe_active, PE_Like=pe_like, PA_Sick=pa_sick,
        Grade_Avg=g_avg, Grade_Math=g_math, Grade_English=g_eng,
        F1=F[0], F2=F[1], F3=F[2], F4=F[3], F5=F[4], Func_Total=func_total,
        Meet_Screen=meet_screen, Meet_Sleep=meet_sleep, Meet_PA=meet_pa, Guidelines_Count=guidelines,
    ))

df = pd.DataFrame(rows)
# shuffle so groups are interleaved, keep IDs sequential
df = df.sample(frac=1, random_state=7).reset_index(drop=True)
df['ID'] = range(1, N + 1)

codebook = [
 ('ID', 'מספר נבדק', '1-65', ''),
 ('Group_LD', 'קבוצה', '1 = עם לקות למידה (קבוצת המחקר), 0 = ללא לקות למידה (קבוצת השוואה)', 'חלק א שאלה 5'),
 ('Gender', 'מין', '1 = בן, 2 = בת', 'חלק א שאלה 1'),
 ('Age', 'גיל', '12-18', 'חלק א שאלה 2'),
 ('Grade', 'כיתה', '7 = ז, 8 = ח, 9 = ט, 10 = י, 11 = יא, 12 = יב', 'חלק א שאלה 3'),
 ('Sector', 'מגזר', '1 = יהודי, 2 = ערבי, 3 = דרוזי', 'חלק א שאלה 4'),
 ('LD_Reading', 'לקות בקריאה', '0 = לא, 1 = כן', 'חלק א שאלה 6'),
 ('LD_Writing', 'לקות בכתיבה', '0 = לא, 1 = כן', 'חלק א שאלה 6'),
 ('LD_Math', 'לקות בחשבון', '0 = לא, 1 = כן', 'חלק א שאלה 6'),
 ('Accommodations', 'התאמות בלמידה או בבחינות', '0 = לא, 1 = כן', 'חלק א שאלה 7'),
 ('ADHD', 'אבחנת הפרעת קשב', '0 = לא, 1 = כן', 'חלק א שאלה 8'),
 ('Medication', 'טיפול תרופתי לקשב', '0 = לא, 1 = כן', 'חלק א שאלה 9'),
 ('TV_WD / TV_WE', 'טלוויזיה וסרטונים, שעות ביום', 'יום חול / סוף שבוע', 'חלק ב טבלה'),
 ('Games_WD / Games_WE', 'משחקים, שעות ביום', 'יום חול / סוף שבוע', 'חלק ב טבלה'),
 ('Social_WD / Social_WE', 'רשתות חברתיות והודעות, שעות ביום', 'יום חול / סוף שבוע', 'חלק ב טבלה'),
 ('Study_WD / Study_WE', 'מסך לצורכי לימודים, שעות ביום', 'יום חול / סוף שבוע (לא נכלל בזמן פנאי)', 'חלק ב טבלה'),
 ('Screen_Leisure_WD', 'זמן מסך פנאי ביום חול', 'TV + Games + Social', 'מחושב'),
 ('Screen_Leisure_WE', 'זמן מסך פנאי בסוף שבוע', 'TV + Games + Social', 'מחושב'),
 ('Screen_Leisure_Avg', 'זמן מסך פנאי יומי ממוצע', '(5 × יום חול + 2 × סוף שבוע) / 7. המשתנה המרכזי של החשיפה למסכים', 'מחושב'),
 ('Screen_Bedroom', 'מסך בחדר השינה בלילה', '0 = לא, 1 = כן', 'חלק ב שאלה 1'),
 ('Screen_BeforeBed', 'מסך בחצי השעה שלפני השינה', '0 = אף פעם, 1 = לעיתים רחוקות, 2 = לפעמים, 3 = כמעט תמיד, 4 = תמיד', 'חלק ב שאלה 2'),
 ('Night_Wake', 'התעוררויות בלילה בגלל הטלפון', '0 = 0, 1 = 1-2, 2 = 3-4, 3 = 5 ומעלה', 'חלק ב שאלה 3'),
 ('Bedtime_WD / Bedtime_WE', 'שעת שינה', 'שעה עשרונית. 23.5 = 23:30. ערך מעל 24 = אחרי חצות (25.0 = 01:00)', 'חלק ג שאלות 1, 3'),
 ('Waketime_WD / Waketime_WE', 'שעת קימה', 'שעה עשרונית. 6.75 = 06:45', 'חלק ג שאלות 2, 4'),
 ('Sleep_Dur_WD', 'משך שינה ביום חול (שעות)', 'Waketime + 24 - Bedtime', 'מחושב'),
 ('Sleep_Dur_WE', 'משך שינה בסוף שבוע (שעות)', 'Waketime + 24 - Bedtime', 'מחושב'),
 ('Weekend_Shift', 'איחור שעת השינה בסוף השבוע (שעות)', 'Bedtime_WE - Bedtime_WD. ערך גבוה = שינה לא סדירה', 'מחושב'),
 ('Sleep_Onset', 'זמן הירדמות', '1 = עד 15 דקות, 2 = 15-30, 3 = 30-60, 4 = יותר משעה', 'חלק ג שאלה 5'),
 ('S1-S8', 'פריטי סולם הישנוניות היומית', '0 = אף פעם ... 4 = תמיד. S3 פריט הפוך (ערנות)', 'חלק ג סולם ישנוניות'),
 ('Sleepiness_Total', 'ציון ישנוניות כולל', 'סכום S1..S8 לאחר היפוך S3 (4 - S3). טווח 0-32', 'מחושב'),
 ('PA_Days60', 'ימים בשבוע האחרון עם 60 דקות פעילות', '0-7', 'חלק ד שאלה 1'),
 ('PA_Freq_Vig', 'תדירות פעילות נמרצת מחוץ לבית הספר', '1 = כל יום, 2 = 4-6 בשבוע, 3 = 2-3 בשבוע, 4 = פעם בשבוע, 5 = פעם בחודש, 6 = פחות מפעם בחודש, 7 = אף פעם (ערך גבוה = פחות פעילות)', 'חלק ד שאלה 2'),
 ('PA_Hours_Vig', 'שעות פעילות נמרצת בשבוע', '1 = אף שעה, 2 = כחצי שעה, 3 = כשעה, 4 = 2-3 שעות, 5 = 4-6 שעות, 6 = 7 ומעלה', 'חלק ד שאלה 3'),
 ('PA_Club', 'חוג או קבוצה', '0 = לא, 1 = כן לא תחרותי, 2 = כן תחרותי', 'חלק ד שאלה 4'),
 ('PE_Active', 'פעילות בשיעורי חינוך גופני', '0 = לא משתתף, 1 = כמעט אף פעם, 2 = לפעמים, 3 = לעיתים קרובות, 4 = תמיד', 'חלק ד שאלה 5'),
 ('PE_Like', 'אהבה לשיעורי חינוך גופני', '1 = כלל לא ... 5 = במידה רבה מאוד', 'חלק ד שאלה 6'),
 ('PA_Sick', 'מחלה או מניעה בשבוע האחרון', '0 = לא, 1 = כן', 'חלק ד שאלה 7'),
 ('Grade_Avg', 'ממוצע ציונים בתעודה האחרונה', '50-100', 'חלק ה שאלה 1'),
 ('Grade_Math', 'ציון אחרון במתמטיקה', '40-100', 'חלק ה שאלה 2'),
 ('Grade_English', 'ציון אחרון באנגלית', '40-100', 'חלק ה שאלה 3'),
 ('F1-F5', 'פריטי תפקוד יומי בלמידה', '1 = כלל לא ... 5 = במידה רבה מאוד. F4 פריט הפוך (עייפות בשיעורים)', 'חלק ה סולם תפקוד'),
 ('Func_Total', 'ציון תפקוד יומי בלמידה', 'ממוצע F1, F2, F3, (6 - F4), F5. טווח 1-5', 'מחושב'),
 ('Meet_Screen', 'עמידה בהמלצת מסכים', '1 = Screen_Leisure_Avg עד 2 שעות', 'מחושב'),
 ('Meet_Sleep', 'עמידה בהמלצת שינה', '1 = 9-11 שעות עד גיל 13, 8-10 שעות מגיל 14 (לפי Sleep_Dur_WD)', 'מחושב'),
 ('Meet_PA', 'עמידה בהמלצת פעילות גופנית', '1 = PA_Days60 שווה 7 (60 דקות בכל יום)', 'מחושב'),
 ('Guidelines_Count', 'מספר ההמלצות שבהן התלמיד עומד', '0-3', 'מחושב'),
]

guide = [
 ('לפני הכול', 'File > Open > הקובץ CSV או XLSX (JASP קורא CSV באופן הכי יציב). לחיצה על כותרת עמודה: לוודא ש-Group_LD, Gender, Sector, PA_Club, Meet_* ו-Guidelines_Count מוגדרים Nominal, וכל השאר Scale.'),
 ('מהימנות הסולמות', 'Reliability > Unidimensional Reliability. להכניס S1..S8 ולסמן את S3 כ-Reverse-Scaled Item. לחזור עם F1..F5 ולסמן F4 כהפוך. לדווח אלפא של קרונבך (ו-McDonald omega).'),
 ('שאלה 1: תיאור הדפוסים', 'Descriptives > Descriptive Statistics. משתנים: Screen_Leisure_Avg, Sleep_Dur_WD, Sleep_Dur_WE, Weekend_Shift, Sleepiness_Total, PA_Days60, Grade_Avg, Func_Total. Split: Group_LD. לסמן Mean, SD, Min, Max. לשיעורי העמידה: Descriptives עם Meet_Screen, Meet_Sleep, Meet_PA, Guidelines_Count וסימון Frequency tables, Split לפי Group_LD.'),
 ('השערה 1: מסכים ושינה', 'Regression > Correlation. משתנים: Screen_Leisure_Avg, Sleep_Dur_WD, Bedtime_WD, Weekend_Shift, Sleepiness_Total, Screen_BeforeBed. Pearson (ו-Spearman אם ההתפלגות לא נורמלית, לבדוק ב-Descriptives > Distribution plots / Shapiro-Wilk). לסנן לקבוצת המחקר בלבד: לחיצה על עמודת Group_LD ולסמן רק את הערך 1.'),
 ('השערה 2: מסכים ופעילות גופנית', 'Regression > Correlation: Screen_Leisure_Avg עם PA_Days60, PA_Hours_Vig, PA_Freq_Vig (שימו לב: PA_Freq_Vig הפוך, ערך גבוה = פחות פעילות).'),
 ('השערה 3: מסכים, שינה, פעילות ותפקוד לימודי', 'Regression > Correlation: Screen_Leisure_Avg, Sleep_Dur_WD, Weekend_Shift, PA_Days60 עם Grade_Avg, Grade_Math, Grade_English, Func_Total. ואז Regression > Linear Regression: Dependent = Grade_Avg (או Func_Total); Covariates = Screen_Leisure_Avg, Sleep_Dur_WD, Weekend_Shift, PA_Days60, Sleepiness_Total; Factors = Group_LD, Gender. לדווח R², Beta, p.'),
 ('השערה 4: מספר ההמלצות ותפקוד', 'ANOVA > ANOVA. Dependent = Grade_Avg; Fixed Factor = Guidelines_Count. לסמן Descriptives ו-Post Hoc (Tukey). אם קבוצת 3 ההמלצות קטנה מדי, לאחד 2 ו-3 לקטגוריה אחת דרך Data > Compute column.'),
 ('השערה 5: השוואה בין הקבוצות', 'T-Tests > Independent Samples T-Test. Grouping = Group_LD. Dependent: Sleep_Dur_WD, Sleepiness_Total, PA_Days60, Screen_Leisure_Avg, Grade_Avg, Func_Total. לסמן Descriptives, Effect size (Cohen d) ו-Assumption checks (Levene). לשיעורי העמידה: Frequencies > Contingency Tables. Rows = Group_LD, Columns = Meet_Screen (ואז Meet_Sleep, Meet_PA). לסמן Chi-squared ו-Row percentages.'),
 ('בקרה על הפרעת קשב', 'לחזור על מבחני t עם Grouping = ADHD בתוך קבוצת המחקר, או להוסיף ADHD כ-Factor ברגרסיה.'),
 ('תרשימים', 'Descriptives > Plots: Boxplots עם Split לפי Group_LD; Correlation > Plots: Scatter plots. אפשר לייצא תרשים בלחיצה ימנית > Copy/Save image.'),
]

wb = Workbook()
hdr_font = Font(name='Arial', bold=True, size=10, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='1F4E78')
body_font = Font(name='Arial', size=10)
thin = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# --- sheet 1: notice ---
ws0 = wb.active; ws0.title = 'הערה'
ws0.sheet_view.rightToLeft = True
notes = [
 'נתונים מדומים לתרגול בלבד',
 'הקובץ נוצר בסימולציה ממוחשבת (Python, זרע אקראי 2026) לפי מבנה השאלון בנספח א של העבודה. אין בו נבדקים אמיתיים.',
 'מטרתו: להכין ולתרגל מראש את הניתוחים ב-JASP (מהימנות, מתאמים, מבחני t, חי בריבוע, ANOVA, רגרסיה) לפני איסוף הנתונים בשטח.',
 'אין להציג את הנתונים האלה או תוצאות שחושבו מהם כממצאי המחקר. לאחר איסוף השאלונים, יש להקליד את התשובות האמיתיות באותו מבנה עמודות (גיליון "נתונים") ולהריץ את אותם ניתוחים.',
 'המדגם המדומה: 65 תלמידים, 40 עם לקות למידה (Group_LD = 1) ו-25 ללא לקות (Group_LD = 0), כיתות ז-יב.',
 'הקשרים שהוכנסו לסימולציה הם בכיוון ההשערות של העבודה (יותר מסכים = פחות שינה ופחות פעילות; לקות למידה = יותר ישנוניות ופחות פעילות), עם רעש אקראי. לכן חלק מההשערות יאוששו וחלק לא, כמו בנתונים אמיתיים.',
 'משתנים מחושבים (Screen_Leisure_*, Sleep_Dur_*, Weekend_Shift, Sleepiness_Total, Func_Total, Meet_*, Guidelines_Count) חושבו לפי הנוסחאות שבגיליון "קודבוק". בנתונים האמיתיים אפשר לחשב אותם באקסל או ב-JASP דרך Compute column.',
 'גיליונות: נתונים | קודבוק | מדריך JASP. קובץ CSV מקביל מצורף לטעינה ישירה ב-JASP.',
]
ws0.column_dimensions['A'].width = 120
for i, t in enumerate(notes, 1):
    c = ws0.cell(row=i, column=1, value=t)
    c.font = Font(name='Arial', size=12 if i == 1 else 10, bold=(i == 1), color=('C00000' if i == 1 else '000000'))
    c.alignment = Alignment(wrap_text=True, horizontal='right', vertical='top', readingOrder=2)
    ws0.row_dimensions[i].height = 22 if i == 1 else 34

# --- sheet 2: data ---
ws = wb.create_sheet('נתונים')
cols = list(df.columns)
for j, name in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=name); c.font = hdr_font; c.fill = hdr_fill
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = border
    ws.column_dimensions[get_column_letter(j)].width = max(9, min(18, len(name) + 2))
for i, row in enumerate(df.itertuples(index=False), 2):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=(float(v) if isinstance(v, (np.floating,)) else int(v) if isinstance(v, (np.integer,)) else v))
        c.font = body_font; c.border = border; c.alignment = Alignment(horizontal='center')
ws.freeze_panes = 'B2'
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{N + 1}"
ws.row_dimensions[1].height = 32

# --- sheet 3: codebook ---
ws2 = wb.create_sheet('קודבוק'); ws2.sheet_view.rightToLeft = True
heads = ['שם המשתנה', 'תיאור', 'ערכים / קידוד', 'מקור בשאלון']
widths = [24, 34, 70, 22]
for j, (h, w) in enumerate(zip(heads, widths), 1):
    c = ws2.cell(row=1, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
    c.alignment = Alignment(horizontal='center', readingOrder=2); ws2.column_dimensions[get_column_letter(j)].width = w
for i, rec in enumerate(codebook, 2):
    for j, v in enumerate(rec, 1):
        c = ws2.cell(row=i, column=j, value=v); c.font = body_font; c.border = border
        c.alignment = Alignment(wrap_text=True, vertical='top', horizontal=('left' if j == 1 else 'right'), readingOrder=(1 if j == 1 else 2))

# --- sheet 4: JASP guide ---
ws3 = wb.create_sheet('מדריך JASP'); ws3.sheet_view.rightToLeft = True
for j, (h, w) in enumerate(zip(['שלב / השערה', 'מה לעשות ב-JASP'], [30, 120]), 1):
    c = ws3.cell(row=1, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = border
    c.alignment = Alignment(horizontal='center', readingOrder=2); ws3.column_dimensions[get_column_letter(j)].width = w
for i, (a, b) in enumerate(guide, 2):
    for j, v in enumerate((a, b), 1):
        c = ws3.cell(row=i, column=j, value=v); c.font = body_font; c.border = border
        c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='right', readingOrder=2)
    ws3.row_dimensions[i].height = 60

wb.save('נתונים_מדומים_לתרגול_JASP_N65.xlsx')
df.to_csv('simulated_data_N65_for_JASP.csv', index=False, encoding='utf-8-sig')

# quick sanity summary
print(df.groupby('Group_LD')[['Screen_Leisure_Avg', 'Sleep_Dur_WD', 'Sleepiness_Total', 'PA_Days60', 'Grade_Avg', 'Func_Total']].mean().round(2))
print('meet rates:', df[['Meet_Screen', 'Meet_Sleep', 'Meet_PA']].mean().round(2).to_dict(), 'guidelines:', df['Guidelines_Count'].value_counts().sort_index().to_dict())
print('corr screen-sleep', round(df.Screen_Leisure_Avg.corr(df.Sleep_Dur_WD), 2), 'screen-PA', round(df.Screen_Leisure_Avg.corr(df.PA_Days60), 2), 'screen-grade', round(df.Screen_Leisure_Avg.corr(df.Grade_Avg), 2))
