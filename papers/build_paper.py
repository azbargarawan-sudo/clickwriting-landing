# -*- coding: utf-8 -*-
import re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------- content ----------
TITLE1 = "צדק מאחה בין נפגע לפוגע:"
TITLE2 = "מנקמה לאיחוי בראי הקרימינולוגיה והיהדות"

COVER = [
    "הקריה האקדמית אונו",
    "הפקולטה למדעי הרוח והחברה",
    "החוג למדעי היהדות, לימודים לתואר מוסמך",
    "עבודה מסכמת: יהדות וקרימינולוגיה",
    "הקורס: ויקטימולוגיה בראייה רב-תרבותית",
    "מרצה: ד\"ר הדר פרנקו גלאור",
]
COVER2 = [
    "מגיש/ה: «שם מלא»",
    "ת.ז.: «מספר תעודת זהות»",
    "תאריך הגשה: «תאריך עברי ולועזי»",
]

INTRO = [
    ("פתיח", None),
    ("עבודת הסמינר ביהדות עסקה ב«נושא עבודת הסמינר ביהדות», ועבודת הסמינר בקרימינולוגיה עסקה ב«נושא עבודת הסמינר בקרימינולוגיה». העבודה הנוכחית עוסקת בנושא אחר: צדק מאחה כתגובה חברתית לפשיעה, ומקומם של הנפגע והפוגע בה.", None),
]

SEC_A = [
    ("סעיף א': הצגת הנושא", None),
    ("צדק מאחה הוא גישה לתגובה חברתית לפשיעה הרואה בעבירה בראש ובראשונה פגיעה בבני אדם וביחסים שביניהם, ולא רק הפרה של חוק המדינה. במקום לשאול איזה חוק הופר ומהו העונש הראוי, הגישה שואלת מי נפגע, מהם צרכיו ומי אחראי לתיקון. ליבת ההליך היא מפגש מונחה בין הנפגע לפוגע, לעיתים בנוכחות תומכים ונציגי קהילה, שבו הפוגע מקבל אחריות על מעשיו והצדדים מגבשים יחד הסכם לתיקון הנזק (קלר-חלמיש ופלג-קוריאט, 2020).", None),
    ("בישראל פועלות בשירות המבחן לנוער שתי תוכניות ממשלתיות: קד\"ם (קבוצות דיון משפחתיות) והנ\"ף (היוועדות נפגע-פוגע). לצידן פועלת תוכנית מקבילה בשירות המבחן למבוגרים. ההליך אינו תחליף להליך הפלילי אלא משולב בו, ותוצאותיו נשקלות בהחלטה על סגירת התיק או בגזר הדין (קלר-חלמיש ופלג-קוריאט, 2020). בשנים האחרונות מתרחב השימוש בגישה גם לפגיעות מיניות, תחום שנחשב במשך שנים בלתי מתאים לה. מחקר ישראלי שראיין שש-עשרה שורדות אלימות מינית שהשתתפו במפגשי צדק מאחה מצא כי אצל רבות מהן התפתחה \"סליחה דיאלוגית\": ירידה ברגשות השליליים כלפי הפוגע שצמחה מתוך השיח במפגש, ולא כדרישה שהוצבה מבחוץ (הדר וגל, 2023ב). במאמר נוסף מאותו מחקר תיארו הנפגעות את התקופה שלאחר ההליך כחוויה של איחוי רגשי, וכותרתו לקוחה מדברי אחת המשתתפות על תחושה שתופרים לה את הלב (הדר וגל, 2023א).", None),
    ("הראיות הבינלאומיות זהירות יותר. מטא-אנליזה מעודכנת של 79 הערכות של תוכניות צדק מאחה לנוער מצאה ירידה קטנה עד בינונית בעבריינות חוזרת, לצד עלייה ניכרת בתחושת ההגינות ובשביעות הרצון של הנפגעים ושל הנערים (Kimbrell et al., 2023). מטא-אנליזה נוספת, שכללה גם תוכניות למבוגרים, מצאה ירידה מובהקת אך קטנה בעבריינות חוזרת כללית, וללא השפעה על עבריינות אלימה חוזרת (Fulham et al., 2025).", None),
    ("ההיבט הרב-תרבותי בולט בישראל במיוחד, והוא ניכר בעמדות הציבור. במחקר ישראלי עדכני תפסו משתתפים יהודים פגיעה מינית בתוך המשפחה כחמורה יותר מפגיעה מחוץ לה ותמכו פחות בצדק מאחה במקרים אלה, ואילו משתתפים ערבים לא הבחינו בין שני סוגי הפגיעה, ותמיכתם בהליך הושפעה פחות מחומרת העבירה (Peleg-Koriat, 2025).", None),
]

SEC_B = [
    ("סעיף ב': ניתוח קרימינולוגי", None),
    ("תיאוריית התיוג. נקודת המוצא של תיאוריית התיוג היא יחסיות הסטייה: אין מעשה סוטה כשלעצמו, אלא מעשה שהחברה הגיבה עליו והגדירה אותו ככזה. התגובה החברתית לפשיעה, ובעיקר התגובה הפורמלית של משטרה ובית משפט, מדביקה לאדם תווית של \"עבריין\", והתווית מעצבת את זהותו העצמית ואת יחסם של אחרים אליו. כאן עוברת ההבחנה בין עבריינות ראשונית, מעשה ראשון שאפשר להסבירו בגורמים מזדמנים, לבין עבריינות משנית, שבה הסטייה נעשית לאורח חיים בעקבות התיוג (Lemert, 1951, כפי שמצוטט אצל Bernburg, 2019). סקירת המחקר העדכני מראה שתיוג פלילי רשמי מגדיל את הסיכוי לעבריינות בהמשך בשני מסלולים: הדרה מהזדמנויות קונבנציונליות בחינוך ובתעסוקה, והשתקעות בקבוצות עברייניות שבהן התווית אינה מכשול אלא יתרון (Bernburg, 2019).", "bold_lead"),
    ("צדק מאחה נולד מתוך ביקורת זו. בריית'ווייט הבחין בין בִּיוּש מתייג, המגנה את האדם ומרחיק אותו, לבין ביוש משלב, המגנה את המעשה בלבד ומסתיים במחוות של קבלה מחדש וסליחה (Braithwaite, 1989, כפי שמצוטט אצל Bernburg, 2019). המפגש המאחה הוא הגלמה מוסדית של הביוש המשלב: הפוגע עומד מול הנזק בנוכחות אנשים שחשובים לו, מקבל אחריות, ומתקבל בחזרה לקהילה בתום ההליך. בקד\"ם המשפחה המורחבת היא הקהילה המביישת והמשלבת גם יחד, ותוצאה מוצלחת מאפשרת סגירת תיק ללא רישום פלילי, כלומר עצירה של המעבר מעבריינות ראשונית למשנית. הממצאים המתונים של המטא-אנליזה מתיישבים עם ההסבר: ההשפעה על עבריינות חוזרת קיימת, אך היא תלויה באיכות המפגש ובמידה שבה הקהילה אכן מקבלת את הנער בחזרה (Kimbrell et al., 2023).", None),
    ("תיאוריית התיוג חלה גם על הנפגע. ואן דייק טוען שבתרבות המערבית התווית \"קורבן\" נושאת קונוטציה של סבל פסיבי וחוסר אונים, ומחקרים פסיכולוגיים-חברתיים עדכניים מאששים כי מי שמתויג כקורבן נתפס כחלש ותלותי (Van Dijk, 2020). ההליך הפלילי מקבע את התווית, שכן הנפגע הוא בו עד ולא צד. צדק מאחה מציע מסלול הפוך: הנפגע בוחר אם להשתתף, מדבר בקולו ומעצב את הסכם התיקון, ובכך נע מתווית של קורבן לעמדה של שורד פעיל.", None),
    ("ויקטימולוגיה פמיניסטית. הגישה השנייה בוחנת את הנושא מכיוון מנוגד. הוויקטימולוגיה הפמיניסטית מדגישה את האופי הממוגדר של הקורבנות, את הקשר בין קורבנות נשים לבין עבריינותן, ואת תפקידה של האלימות כלפי נשים, ואיום האלימות, כאמצעי של פיקוח חברתי בלתי פורמלי (Clay-Warner & Edgemon, 2020). מנקודת מבט זו ההליך הפלילי בפגיעות מיניות מייצר קורבנות משנית: חקירה נגדית, מיתוסים של אונס ואשמת הקורבן. אך גם צדק מאחה חשוד בעיניה, משום שהוא עלול לשחזר במפגש את פערי הכוח שבין הפוגע לנפגעת, להפעיל עליה לחץ לסלוח ולהפריט פגיעה שהיא בעיה ציבורית (קלר-חלמיש ופלג-קוריאט, 2020).", "bold_lead"),
    ("מחאת #MeToo שינתה את המאזן. קלר-חלמיש ופלג-קוריאט (2020) טוענות כי המחאה חשפה את היקף הפגיעות ואת כישלונו של ההליך הפלילי לתת מענה לרובן, ולכן יש להנגיש צדק מאחה לפגיעות מיניות בתנאים מגוננים: השתתפות וולונטרית, הכנה ממושכת, מנחים מיומנים וזכותה של הנפגעת לעצור בכל שלב. ממצאי הראיונות עם השורדות תומכים בעמדה זו. הסליחה, כשהופיעה, לא הייתה תנאי ולא דרישה אלא תוצר של הדיאלוג, ומשתתפות שלא סלחו דיווחו אף הן על הכרה ועל השבת השליטה (הדר וגל, 2023ב). התגובה החברתית לנפגע עצמו ממוגדרת אף היא. ניתוח של 1,890 תגובות גולשים לחשיפות של נפגעי פגיעה מינית במסגרת #MeToo בישראל מצא שנשים זכו לתמיכה רגשית ורשתית רבה יותר מגברים, ואילו גברים קיבלו בעיקר תמיכה הערכתית (לוינשטיין-ברקאי, 2020). מי שאינו תואם את דמות הקורבן המצופה זוכה לפחות הכרה. הביקורת הפמיניסטית על \"הקורבן האידיאלי\", זו שזוכה להכרה רק אם היא תמימה וחסרת אשמה, מהדהדת גם בממצאים על הבדלים אתניים בתמיכה בצדק מאחה: מי נחשב ראוי להליך מאחה ומי לענישה בלבד תלוי בזהות המעורבים, ולא רק במעשה (Peleg-Koriat, 2025).", None),
    ("ממצא זה מתחדד במחקר ניסויי נוסף שנערך בישראל. בשני ניסויים, האחד עם 446 משתתפים והשני עם 560, נמצאה הטיה אתנית מובהקת בתמיכה בצדק מאחה בהתאם לזהותם של הפוגע, הנפגע והמשיב, והטיה זו הייתה חזקה יותר בהקשרים פוליטיים טעונים מאשר בעבירות פליליות רגילות (Peleg-Koriat et al., 2025). התיווך נעשה דרך אמונות בדבר יכולתו של האדם להשתנות: מי שמאמין שהפוגע יכול להשתנות תומך בהליך, והאמונה הזו עצמה תלויה בזהות. מנקודת מבטה של תיאוריית התיוג זו הדגמה ישירה של יחסיות הסטייה. אותו מעשה נתפס כניתן לתיקון או כבלתי ניתן לתיקון לפי מי שביצע אותו, וההחלטה מי יופנה להליך מאחה ומי לענישה היא עצמה תגובה חברתית שמייצרת עבריינים.", None),
    ("שתי התיאוריות משלימות זו את זו ומתנגשות בנקודה אחת. התיוג ממקד את המבט בפוגע ובשאלה כיצד למנוע את קיבוע זהותו העבריינית. הוויקטימולוגיה הפמיניסטית ממקדת אותו בנפגעת ובפערי הכוח. הצלחת ההליך תלויה ביכולת להחזיק את שתיהן בו-זמנית: שילוב מחדש של הפוגע, בלי שהמחיר יהיה קולה של הנפגעת.", None),
]

SEC_C = [
    ("סעיף ג': ניתוח מפרספקטיבה יהודית", None),
    ("הפרספקטיבה היהודית על צדק מאחה נלמדה בתואר בשני מוקדים: המקרא וספרות חז\"ל בקורס «שם הקורס», וספרות ההלכה של ימי הביניים ואילך בקורס «שם הקורס». הערכים המרכזיים שעלו בהם ביחס לנושא הם וידוי, השבה, פיוס ומחילה.", None),
    ("מקורות התשתית. התורה מכירה בשני מעגלי פגיעה ומטילה חובות על שני הצדדים. על הפוגע: \"והתודו את חטאתם אשר עשו והשיב את אשמו בראשו וחמישיתו יוסף עליו ונתן לאשר אשם לו\" (במדבר ה, ז). כלומר וידוי, השבה בתוספת חומש, ומסירה ישירה לידי הנפגע ולא לקופת הציבור. על הנפגע: \"לא תשנא את אחיך בלבבך הוכח תוכיח את עמיתך... לא תקם ולא תטר... ואהבת לרעך כמוך\" (ויקרא יט, יז-יח). האיסור על נקמה ונטירה אינו מותיר את הנפגע בשתיקה, אלא מחליף את הנקמה בתוכחה, כלומר בפנייה ישירה אל הפוגע.", "bold_lead"),
    ("חז\"ל הפכו עקרונות אלו למנגנון. המשנה קובעת: \"עבירות שבין אדם למקום יום הכיפורים מכפר, עבירות שבין אדם לחברו אין יום הכיפורים מכפר עד שירצה את חברו\" (משנה, יומא ח, ט). המבנה מבחין בין הפגיעה בסדר הציבורי, שהמערכת יכולה לכפר עליה, לבין הפגיעה באדם, שאיש אינו רשאי למחול עליה בשמו. במסכת בבא קמא הכלל מפורש עוד יותר: החובל ששילם \"אין נמחל לו עד שיבקש ממנו\", והנפגע מוזהר מצידו \"שלא יהא המוחל אכזרי\" (משנה, בבא קמא ח, ז; בבלי, בבא קמא צב ע\"א). הגמרא לומדת את חובת המחילה מאברהם, שהתפלל לרפואת אבימלך לאחר שזה לקח את שרה (בראשית כ, יז), ומוסיפה: \"כל המבקש רחמים על חברו והוא צריך לאותו דבר, הוא נענה תחילה\" (בבלי, בבא קמא צב ע\"א). הנפגע שפועל לטובת הפוגע אינו מוותר על עצמו אלא נרפא ראשון. התשלום הכספי הוא רק תחילת התיקון. סופו במפגש.", None),
    ("מקורות ממשיכים. הרמב\"ם קודד את ההלכה ופירט את ההליך. עבירות שבין אדם לחברו \"אינו נמחל לו לעולם עד שיתן לחברו מה שהוא חייב לו וירצהו\", ואם הנפגע מסרב, הפוגע \"מביא לו שורה של שלושה בני אדם מרעיו ופוגעין בו ומבקשין ממנו\", עד שלוש פעמים. אם גם אז הנפגע מסרב, הפוגע פטור, \"וזה שלא מחל הוא החוטא\" (רמב\"ם, הלכות תשובה ב, ט). ההליך ההלכתי מזכיר בפרטיו את המפגש המאחה: פנייה ישירה, מלווים מטעם הפוגע, חזרה על הניסיון, וסיום בהכרעה של הנפגע. בהלכה שלאחריה מוטלת חובה על הנפגע: \"אסור לאדם להיות אכזרי ולא יתפייס, אלא יהא נוח לרצות וקשה לכעוס\" (רמב\"ם, הלכות תשובה ב, י). בהלכות חובל ומזיק מוסיף הרמב\"ם שפגיעה בגוף שונה מפגיעה בממון: גם לאחר התשלום \"אינו מתכפר לו עד שיבקש מן הנחבל וימחול לו\" (רמב\"ם, הלכות חובל ומזיק ה, ט). השולחן ערוך קבע את הדברים הלכה למעשה כחלק מההכנה ליום הכיפורים (שולחן ערוך, אורח חיים תרו, א).", "bold_lead"),
    ("התמונה העולה מן המקורות ברורה. היהדות אינה מציעה צדק מאחה במקום ענישה, שהרי דיני החובל והגזלן נשארים בתוקפם, אלא קובעת שהענישה והפיצוי אינם סוף הסיפור. הכפרה השלמה מותנית בשלושה מרכיבים שמערכת המשפט המודרנית אינה דורשת: וידוי אישי, בקשת מחילה ישירה מן הנפגע, ומחילה של הנפגע עצמו. בד בבד היא מטילה על הנפגע אחריות מוסרית משלו: לא לנקום, לא לנטור, ולא להיות אכזרי.", None),
]

SEC_D = [
    ("סעיף ד': אינטגרציה בין פרספקטיבות", None),
    ("הניתוח הקרימינולוגי והפרספקטיבה היהודית משלימים זה את זה ברובד המבני ומתנגשים ברובד הנורמטיבי. ברובד המבני הדמיון בולט: ההבחנה של חז\"ל בין עבירות שבין אדם למקום לעבירות שבין אדם לחברו מקבילה להבחנה של צדק מאחה בין הפגיעה במדינה לפגיעה באדם; \"שורה של שלושה בני אדם מרעיו\" מקבילה לתומכים במפגש הנפגע-פוגע; ודרישת הרמב\"ם מבעל התשובה לומר \"אני אחר ואיני אותו האיש שעשה אותן המעשים\" (רמב\"ם, הלכות תשובה ב, ד) היא ניסוח מדויק של המטרה שתיאוריית התיוג מציבה: זהות חדשה במקום עבריינות משנית.", None),
    ("ההבדל הראשון הוא במעמד הטיעון. הקרימינולוגיה תיאורית ואמפירית, והיא מודדת אם ההליך מפחית עבריינות חוזרת ומיטיב עם הנפגע (Kimbrell et al., 2023). היהדות נורמטיבית: היא אינה שואלת אם המפגש עובד אלא קובעת שהוא חובה. ההבדל השני, החריף יותר, נוגע לנפגע. הוויקטימולוגיה הפמיניסטית, וכמוה החשיבה הטיפולית העכשווית, עומדות על כך שסליחה אינה יכולה להיות חובה ואינה יכולה לבוא מבחוץ, אלא רק בקצב של הנפגע (הדר וגל, 2023ב; Clay-Warner & Edgemon, 2020). ההלכה, לעומת זאת, מכנה סרבן מחילה \"אכזרי\". הסתירה מתמתנת כשמבחינים בין הליך לעמדה מוסרית: ההלכה אינה כופה מפגש, מתירה לפוגע להניח לנפגע לאחר שלוש פניות, ומכוונת את תביעתה אל תודעת הנפגע ולא אל בית הדין. תרומת הקרימינולוגיה היא הזהירות, כלומר ידע אמפירי על פערי כוח ועל תנאי ההצלחה של המפגש. תרומת היהדות היא התביעה המוסרית: הפוגע חייב לעמוד מול מי שפגע בו, והנפגע אינו רק מושא להגנה אלא סובייקט מוסרי בעל אחריות.", None),
]

BIB_HE = [
    "הדר, נ' וגל, ט' (2023א). \"הרגשתי איך תופרים לי את הלב\": החוויה הרגשית של נפגעות אלימות מינית בעקבות תהליכי צדק מאחה. חברה ורווחה, מג(3), 355-375. https://www.gov.il/BlobFolder/reports/molsa-43-3-nataly-h/he/SocialAndWelfareMagazine_Magazine-43-3_43-3-NATALY-H.pdf",
    "הדר, נ' וגל, ט' (2023ב). לסלוח או לא לסלוח: סליחה דיאלוגית על פגיעות מיניות בתהליכי צדק מאחה. קרימינולוגיה ישראלית, יא, 87-116. https://cris.haifa.ac.il/en/publications/לסלוח-או-לא-לסלוח-סליחה-דיאלוגית-על-פגיעות-מיניות-בתהליכי-צדק-מאח",
    "לוינשטיין-ברקאי, ה' (2020). המגדר כן קובע? תמיכה חברתית בקורבנות-גברים של פגיעה מינית לעומת תמיכה בקורבנות-נשים ברשתות חברתיות בישראל. מגמות, נה(2), 167-190. https://info2011.szold.org.il/PDF-Articles/megamot/113341.pdf",
    "קלר-חלמיש, כ' ופלג-קוריאט, ע' (2020). ממחאה לאיחוי: מחאת #MeToo כרקע לפיתוח והנגשה של תהליכי צדק מאחה במקרים של פגיעה מינית. קרימינולוגיה ישראלית, ט, 33-54. https://papers.ssrn.com/sol3/papers.cfm?abstract_id=3901353",
]
BIB_EN = [
    "Bernburg, J. G. (2019). Labeling theory. In M. D. Krohn, N. Hendrix, G. Penly Hall, & A. J. Lizotte (Eds.), Handbook on crime and deviance (2nd ed., pp. 179-196). Springer. https://www.researchgate.net/publication/336312509_LABELING_THEORY",
    "Clay-Warner, J., & Edgemon, T. G. (2020). Feminist approaches to victimology. In S. Walklate, K. Fitz-Gibbon, J. Maher, & J. McCulloch (Eds.), The Emerald handbook of feminism, criminology and social change (pp. 35-50). Emerald Publishing. https://research.birmingham.ac.uk/en/publications/feminist-approaches-to-victimology/",
    "Fulham, L., Blais, J., Rugge, T., & Schultheis, E. A. (2025). The effectiveness of restorative justice programs: A meta-analysis of recidivism and other relevant outcomes. Criminology & Criminal Justice, 25(5), 1486-1512. https://doi.org/10.1177/17488958231215228",
    "Kimbrell, C. S., Wilson, D. B., & Olaghere, A. (2023). Restorative justice programs and practices in juvenile justice: An updated systematic review and meta-analysis for effectiveness. Criminology & Public Policy, 22(1), 161-195. https://doi.org/10.1111/1745-9133.12613",
    "Peleg-Koriat, I. (2025). Ethnic differences in support for restorative justice in sexual abuse cases: The role of offense severity and malleability beliefs. Violence Against Women. Advance online publication. https://pmc.ncbi.nlm.nih.gov/articles/PMC13031375/",
    "Peleg-Koriat, I., Weimann-Saks, D., Asraf, K., & Halperin, E. (2025). Ethnicity and support for restorative justice: The mediating role of malleability beliefs and attribution bias. Journal of Community & Applied Social Psychology, 35, Article e70117. https://doi.org/10.1002/casp.70117",
    "Van Dijk, J. (2020). Victim labeling theory: A reappraisal. In J. Joseph & S. Jergenson (Eds.), An international perspective on contemporary developments in victimology: A festschrift in honor of Marc Groenhuijsen (pp. 73-90). Springer. https://www.researchgate.net/publication/342546073_Victim_Labeling_Theory_A_Reappraisal",
]
BIB_PRIMARY_TITLE = "מקורות יהודיים ראשוניים (מצוטטים בגוף העבודה לפי המספור המקובל)"
BIB_PRIMARY = [
    "בראשית כ, יז. https://www.sefaria.org.il/Genesis.20.17",
    "ויקרא יט, יז-יח. https://www.sefaria.org.il/Leviticus.19.17",
    "במדבר ה, ז. https://www.sefaria.org.il/Numbers.5.7",
    "משנה, יומא ח, ט. https://www.sefaria.org.il/Mishnah_Yoma.8.9",
    "משנה, בבא קמא ח, ז. https://www.sefaria.org.il/Mishnah_Bava_Kamma.8.7",
    "תלמוד בבלי, בבא קמא צב ע\"א. https://www.sefaria.org.il/Bava_Kamma.92a",
    "רמב\"ם, משנה תורה, הלכות תשובה ב, ד; ב, ט-י. https://www.sefaria.org.il/Mishneh_Torah,_Repentance.2",
    "רמב\"ם, משנה תורה, הלכות חובל ומזיק ה, ט. https://www.sefaria.org.il/Mishneh_Torah,_One_Who_Injures_a_Person_or_Property.5.9",
    "שולחן ערוך, אורח חיים, סימן תרו, סעיף א. https://www.sefaria.org.il/Shulchan_Arukh,_Orach_Chayim.606.1",
]

# ---------- helpers ----------
FONT = "David"

def set_rtl(p, rtl=True):
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1' if rtl else '0')
    pPr.append(bidi)

def style_run(r, size=12, bold=False, italic=False, highlight=False, font=FONT):
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts'); rPr.append(rFonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rFonts.set(qn(a), font)
    szCs = OxmlElement('w:szCs'); szCs.set(qn('w:val'), str(size*2)); rPr.append(szCs)
    if bold:
        bCs = OxmlElement('w:bCs'); rPr.append(bCs)
    if italic:
        iCs = OxmlElement('w:iCs'); rPr.append(iCs)
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW

def add_text(p, text, size=12, bold=False, italic=False):
    # «...» marks a placeholder to highlight
    parts = re.split(r'(«[^»]*»)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('«'):
            r = p.add_run('[' + part[1:-1] + ']')
            style_run(r, size, bold, italic, highlight=True)
        else:
            r = p.add_run(part)
            style_run(r, size, bold, italic)

def para(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, rtl=True, size=12, bold=False,
         italic=False, spacing=1.5, after=6, before=0, bold_lead=False, hanging=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    if hanging:
        pf.left_indent = Cm(1.25)
        pf.first_line_indent = Cm(-1.25)
    set_rtl(p, rtl)
    if bold_lead:
        lead, rest = text.split('.', 1)
        add_text(p, lead + '.', size, bold=True)
        add_text(p, rest, size)
    else:
        add_text(p, text, size, bold, italic)
    return p

def heading(doc, text, size=14, page_break_before=False):
    p = para(doc, text, align=WD_ALIGN_PARAGRAPH.RIGHT, size=size, bold=True, after=6, before=12)
    if page_break_before:
        p.paragraph_format.page_break_before = True
    return p

def section(doc, items, page_break=True):
    first = True
    for text, kind in items:
        if first:
            heading(doc, text, page_break_before=page_break)
            first = False
        else:
            para(doc, text, bold_lead=(kind == "bold_lead"))

# ---------- build ----------
doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
for side in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
    setattr(sec, side, Cm(2.54))

# default style
st = doc.styles['Normal']
st.font.name = FONT; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:cs'), FONT)
st.element.rPr.rFonts.set(qn('w:ascii'), FONT)
st.element.rPr.rFonts.set(qn('w:hAnsi'), FONT)

# cover
for i, line in enumerate(COVER):
    para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, size=13 if i == 0 else 12, bold=(i == 0), after=2)
for _ in range(5):
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para(doc, TITLE1, align=WD_ALIGN_PARAGRAPH.CENTER, size=18, bold=True, after=2)
para(doc, TITLE2, align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, after=2)
for _ in range(7):
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
for line in COVER2:
    para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)

# body
section(doc, INTRO, page_break=True)
section(doc, SEC_A, page_break=False)
section(doc, SEC_B, page_break=False)
section(doc, SEC_C, page_break=False)
section(doc, SEC_D, page_break=False)

# bibliography
heading(doc, "רשימת מקורות", page_break_before=True)
for item in BIB_HE:
    para(doc, item, align=WD_ALIGN_PARAGRAPH.RIGHT, spacing=1.5, hanging=True, after=8)
for item in BIB_EN:
    para(doc, item, align=WD_ALIGN_PARAGRAPH.LEFT, rtl=False, spacing=1.5, hanging=True, after=8)
para(doc, BIB_PRIMARY_TITLE, align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, before=10, after=4)
for item in BIB_PRIMARY:
    para(doc, item, align=WD_ALIGN_PARAGRAPH.RIGHT, spacing=1.5, hanging=True, after=6)

out = sys.argv[1]
doc.save(out)

# ---------- word counts ----------
def wc(items):
    return sum(len(t.split()) for t, _ in items[1:])
counts = {"פתיח": wc(INTRO), "א": wc(SEC_A), "ב": wc(SEC_B), "ג": wc(SEC_C), "ד": wc(SEC_D)}
counts["סה\"כ"] = sum(counts.values())
print(counts)

# dump plain text for review
with open(out.replace('.docx', '.md'), 'w', encoding='utf-8') as f:
    f.write(f"# {TITLE1} {TITLE2}\n\n")
    for block in (INTRO, SEC_A, SEC_B, SEC_C, SEC_D):
        f.write(f"## {block[0][0]}\n\n")
        for t, _ in block[1:]:
            f.write(t + "\n\n")
    f.write("## רשימת מקורות\n\n")
    for i in BIB_HE + BIB_EN:
        f.write(i + "\n\n")
    f.write(f"### {BIB_PRIMARY_TITLE}\n\n")
    for i in BIB_PRIMARY:
        f.write(i + "\n\n")
